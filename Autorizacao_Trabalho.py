from tkinter import messagebox, Tk, Toplevel, ttk
import pandas as pd
import datetime
from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time
from ui_helpers import create_loading_dialog

class trabalho():
    try:
        def funcao_a_executar3(self):
            # Reutiliza a janela principal do aplicativo para o progresso.
            root = self.janela

            # Dimensões da janela de loading
            largura = 400
            altura = 100

            # Calcula posição central
            x = (root.winfo_screenwidth() // 2) - (largura // 2)
            y = (root.winfo_screenheight() // 2) - (altura // 2)

            loading, progresso = create_loading_dialog(root)
            # Use a função read_excel para ler o arquivo
            tabela = pd.read_excel('ADITAMENTO.xlsx')

            # cria os arrays
            NOME = []
            CTPS = []
            SERIE_CTPS = []
            CPF = []
            Descricao = []
            Nacionalidades = []
            Funcao = []

            # Busca a data atual do sistema e converte no formato brasileiro
            data_atual = datetime.date.today()
            data = data_atual.strftime('%d/%m/%Y')

            # Iterar sobre as linhas
            for index, linha in tabela.iterrows():
                NOME.append(linha['Nome'])
                CTPS.append(linha['CTPS'])
                SERIE_CTPS.append(linha['Série CTPS'])
                CPF.append(linha['CPF'])
                Descricao.append(linha['Descrição'])
                Nacionalidades.append(linha['Nacionalidade'])
                Funcao.append(linha['Função'])

                # Conta a quantidade de linhas
                quant = len(NOME)
                valores = []

            progresso.configure(maximum=max(quant, 1))

            # Executa um loop gravando as informações no texto
            for dado in range(quant):
                arquivo = f"""Pelo presente instrumento particular, de um lado, SUPER MERCADO CASTANHA LTDA, inscrita no CNPJ/MF sob nº63.082.721/0001-08, doravante denominada EMPREGADORA, e, de outro lado, {NOME[dado]}, {Nacionalidades[dado]}, {Descricao[dado]}, {Funcao[dado]}, portador(a) do CPF nº{CPF[dado]}, e da CTPS nº{CTPS[dado]},  Série {SERIE_CTPS[dado]},  doravante denominado(a) FUNCIONÁRIO(A), têm entre si justo e avençado o presente instrumento de autorização que se regerá pelas cláusulas seguintes e pelas condições descritas no objeto.
        CLÁUSULA PRIMEIRA - OBJETO
        1.1- O presente instrumento contrato tem por OBJETO a autorização expressa do(a) FUNCIONÁRIO(A) para prestar serviços à EMPREGADORA nos dias 24 de dezembro de 2023 e 31 de dezembro de 2023, em conformidade com o aditamento da Convenção Coletiva dos Comerciários da Capital/SINCOVAGA 2023/2023 e a legislação trabalhista vigente.
        1.2- Esta autorização fundamenta-se na inserção do parágrafo 4º à cláusula 44 da referida Convenção Coletiva de Trabalho (CCT), que trata do trabalho aos domingos e feriados.
        CLÁUSULA SEGUNDA - TERMOS E CONDIÇÕES
        2.1- O FUNCIONÁRIO(A) declara, por meio deste instrumento, que AUTORIZA EXPRESSAMENTE E CONCORDA em trabalhar nos dias 24 e 31 de dezembro de 2023, nos termos da CCT e legislação aplicável, recebendo as devidas contraprestações e benefícios previstos na legislação e convenção supracitadas.
        2.2- O(A) funcionário(a) declara que a presente autorização é dada de forma livre e espontânea, sem qualquer coação ou pressão por parte da EMPREGADORA.
        2.3- A EMPREGADORA se compromete a assegurar que toda a remuneração e adicionais relacionados ao trabalho prestado nestas datas especiais sejam pagos de acordo com o que estipula a CCT e a legislação trabalhista vigente.
        2.4- O(A) FUNCIONÁRIO(A) está ciente que o não comparecimento nos dias especificados neste instrumento, serão aplicadas pela EMPREGADORA as penalidades cabíveis na legislação, salvo se ocorrer condições previstas em lei e devidamente comprovadas pelo(a) FUNCIONÁRIO(A).
        2.5- Esta autorização terá vigência específica para os dias 24 e 31 de dezembro de 2023, sem prejuízo das demais cláusulas do contrato de trabalho em vigor entre as partes.
        E, perante a concordância expressa do(a) FUNCIONÁRIO(A), celebra o presente instrumento de autorização, apondo sua assinatura para que surta os devidos efeitos legais e jurídicos.
        São Paulo, {data}.
        ________________________________
        SUPER MERCADO CASTANHA LTDA
        EMPREGADORA
        ________________________________
        FUNCIONÁRIO(A)
        {NOME[dado]}
         CPF: {CPF[dado]}"""

                # Adicione a string ao array
                valores.append(arquivo)

                # Crie um novo documento do Word
                doc = Document()

                # Adicione um cabeçalho
                header = doc.sections[0].header
                paragraph = header.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(
                    'AUTORIZAÇÃO DE TRABALHO EM DIAS ESPECÍFICOS CONFORME ADITAMENTO DA CCT DOS COMERCIÁRIOS DA CAPITAL/SINCOVAGA 2023/2023')

                # Adiciona um texto em negrito
                font = run.font
                font.bold = True
                font.size = Pt(12)

                # #Adicionando imagem ao rodapé
                section = doc.sections[0]
                footer = section.footer
                footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                footer.paragraphs[0].add_run().add_picture('Logo Castanha2.png', width=Inches(1.0), height=Inches(0.5))

                # Adicione as strings ao documento
                for texto in valores:
                    paragrafo = doc.add_paragraph(texto)
                    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    print('ok')
                    # Salve o documento
                    doc.save('ADITAMENTO.docx')
                    # Atualiza barra de progresso
                    progresso['value'] = dado + 1
                    loading.update()
                    time.sleep(0.1)

            loading.destroy()
            messagebox.showwarning(title="Assiduidades", message="Arquivo gerado com sucesso")
    except Exception as erro:
        messagebox.ERROR(f"{erro}")
