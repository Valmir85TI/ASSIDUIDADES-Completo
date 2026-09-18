from tkinter import messagebox
import pandas as pd
import datetime
from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import workbook

class natal():
    try:
        def funcao_a_executar6(self):
            # Use a função read_excel para ler o arquivo
            tabela = pd.read_excel('VALENATAL.xlsx')

            # cria os arrays
            NOME = []
            # CTPS = []
            # SERIE_CTPS = []
            #CPF = []

            # Busca a data atual do sistema e converte no formato brasileiro
            data_atual = datetime.date.today()
            data = data_atual.strftime('%d/%m/%Y')

            # Iterar sobre as linhas
            for index, linha in tabela.iterrows():
                NOME.append(linha['Nome'])
                # CTPS.append(linha['CTPS'])
                # SERIE_CTPS.append(linha['Série CTPS'])
                #CPF.append(linha['CPF'])
                # coluna_local = linha['Descrição do Local']

                # Conta a quantidade de linhas
                quant = len(NOME)
                valores = []

            # Executa um loop gravando as informações no texto
            for dado in range(quant):
                arquivo = f"""     
              Caro colaborador(a) {NOME[dado]}
              Hoje, queremos expressar nossa imensa gratidão pelo seu comprometimento, saiba
              que seu esforço e dedicação são fundamentais para o sucesso da nossa empresa.
              Queremos aproveitar este momento para reconhecer e agradecer pela sua dedicação 
              este presente é uma pequena forma de demonstrar nossa gratidão por tudo."""

                # Adicione a string ao array
                valores.append(arquivo)

                # Crie um novo documento do Word
                doc = Document()

                # # Adicione um cabeçalho
                # header = doc.sections[0].header
                # paragraph = header.paragraphs[0]
                # run = paragraph.add_run(
                #     '                                 TERMO DE CONCESSÃO DO VALE NATALINO\n              Parágrafo Único da Cláusula 65ª da Convenção Coletiva 2024/2025\n')
                #
                # # Adiciona um texto em negrito
                # font = run.font
                # font.bold = True
                # font.size = Pt(12)

                # Adicionando imagem ao rodapé
                section = doc.sections[0]
                footer = section.footer
                footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                footer.paragraphs[0].add_run().add_picture('Logo Castanha2.png', width=Inches(1.5), height=Inches(1.0))

                # Adicione as strings ao documento
                for texto in valores:
                    doc.add_paragraph(texto)
                    print('ok')
                    # Salve o documento
                    doc.save('VALENATAL.docx')
            messagebox.showwarning(title="Assiduidades", message="Arquivo gerado com sucesso")
    except Exception as erro:
        messagebox.ERROR(f"{erro}")
