from tkinter import messagebox, Tk, Toplevel, ttk
import pandas as pd
import datetime
from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import workbook
import time
from ui_helpers import create_loading_dialog

class natalCastanhao():
    try:
        def funcao_a_executar5(self):
            # Reutiliza a janela principal do aplicativo para o progresso.
            root = self.janela

            # Dimensões da janela de loading
            largura = 400
            altura = 100

            # Calcula posição central
            x = (root.winfo_screenwidth() // 2) - (largura // 2)
            y = (root.winfo_screenheight() // 2) - (altura // 2)

            loading, progresso = create_loading_dialog(root)

            # Use a função read_excel para ler o arquivo
            tabela = pd.read_excel('VALENATAL_CASTANHAO.xlsx')

            # cria os arrays
            NOME = []
            # CTPS = []
            # SERIE_CTPS = []
            CPF = []

            # Busca a data atual do sistema e converte no formato brasileiro
            data_atual = datetime.date.today()
            data = data_atual.strftime('%d/%m/%Y')

            # Iterar sobre as linhas
            for index, linha in tabela.iterrows():
                NOME.append(linha['Nome'])
                # CTPS.append(linha['CTPS'])
                # SERIE_CTPS.append(linha['Série CTPS'])
                CPF.append(linha['CPF'])
                # coluna_local = linha['Descrição do Local']

                # Conta a quantidade de linhas
                quant = len(NOME)
                valores = []

            progresso.configure(maximum=max(quant, 1))

            # Executa um loop gravando as informações no texto
            for dado in range(quant):
                arquivo = f"""      Castanhão Atacado Comercio e Distribuidora LTDA., inscrita no CNPJ/MF: 24.274.152/0001-40, empregador, através do presente termo faz registrar a concessão do VALE NATALINO na forma prevista pelo parágrafo único da cláusula 65ª da CCT 2024/2025, mediante os critérios abaixo especificados:

        Conforme previsto no parágrafo único da cláusula 65ª da CCT 2024/2025, a empresa poderá substituir a Cesta Natalina pelo documento-refeição como vale natalino com o valor correspondente aos produtos contemplados na cesta de natal.
        Conforme apurado, os produtos contemplados na cesta de natal equivale ao valor de R$105,00 (cento e cinco reais).
        Todos os funcionários serão contemplados com o recebimento do vale compra (documento-refeição) no valor de R$105,00 (cento e cinco reais) a título de vale natalino.

        O vale compra (documento-refeição) a título de vale natalino no valor de R$105 (cento e cinco reais) somente poderá ser utilizado para aquisição de produtos comercializados pelo empregador.

        Diante do exposto, consensualmente fica celebrado:
        O funcionário declara que está recebendo o vale compra (documento-refeição) no valor de R$105 (cento e cinco reais) a título de vale natalino em substituição à Cesta Natalina conforme previsto no parágrafo único da cláusula 65ª da CCT 2024/2025.

        O funcionário declara que o vale compra (documento-refeição) a título de vale natalino ora recebido não guarda qualquer relação com a remuneração ou com salário in natura e não se aplica os arts. 457 e artigo 458 da CLT. Declara está ciente de que o valor recebido é verba de natureza 100% indenizatória e não comporá a base de cálculo dos tributos e nem do FGTS, assim como declara estar ciente de que o respectivo valor não gera qualquer hipótese de direito adquirido.
        O presente instrumento foi redigido em conformidade com a Lei 13.709/18 Lei Geral de Proteção de Dados Pessoais (LGPD).

        Eu, {NOME[dado]}, CPF: {CPF[dado]}, declaro estar ciente e de pleno acordo com as condições acima estão transcritas, e que consumirei o valor do vale compra (documento-refeição) a título de vale natalino na aquisição de produtos comercializados pelo empregador Castanhão Atacado Comercio e Distribuidora LTDA.
        São Paulo, {data}.
            ________________________________
            Funcionário.
            {NOME[dado]}\n"""

                # Adicione a string ao array
                valores.append(arquivo)

                # Crie um novo documento do Word
                doc = Document()

                # Adicione um cabeçalho
                header = doc.sections[0].header
                paragraph = header.paragraphs[0]
                run = paragraph.add_run(
                    '                                 TERMO DE CONCESSÃO DO VALE NATALINO\n              Parágrafo Único da Cláusula 65ª da Convenção Coletiva 2024/2025\n')

                # Adiciona um texto em negrito
                font = run.font
                font.bold = True
                font.size = Pt(12)

                # Adicionando imagem ao rodapé
                section = doc.sections[0]
                footer = section.footer
                footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                footer.paragraphs[0].add_run().add_picture('Logo Castanhao.jpg', width=Inches(1.5), height=Inches(1.0))

                # Adicione as strings ao documento
                for texto in valores:
                    doc.add_paragraph(texto)
                    print('ok')
                    # Salve o documento
                    doc.save('VALENATAL_CASTANHAO.docx')
                    # Atualiza barra de progresso
                    progresso['value'] = dado + 1
                    loading.update()
                    time.sleep(0.1)

            loading.destroy()
            messagebox.showwarning(title="Assiduidades", message="Arquivo gerado com sucesso")
    except Exception as erro:
        messagebox.ERROR(f"{erro}")
